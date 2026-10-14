import sys
import subprocess

# Ensure python-pptx is installed
try:
    import pptx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    import pptx

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Colors
    DARK_BG = RGBColor(15, 23, 42)      # #0F172A
    CARD_BG = RGBColor(30, 41, 59)      # #1E293B
    ACCENT_TEAL = RGBColor(20, 184, 166) # #14B8A6
    ACCENT_BLUE = RGBColor(59, 130, 246) # #3B82F6
    TEXT_LIGHT = RGBColor(241, 245, 249) # #F1F5F9
    TEXT_MUTED = RGBColor(148, 163, 184)# #94A3B8
    WHITE = RGBColor(255, 255, 255)
    BORDER_COLOR = RGBColor(51, 65, 85)

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="SECHAT ARCHITECTURE"):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_TEAL

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, DARK_BG)

    # Accent decorative box
    dec_shape = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.2), Inches(0.15), Inches(2.8))
    dec_shape.fill.solid()
    dec_shape.fill.fore_color.rgb = ACCENT_TEAL
    dec_shape.line.fill.background()

    # Title Box
    tbox = slide1.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(11), Inches(3))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "SeChat Real-Time Messaging Platform"
    p0.font.size = Pt(38)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_LIGHT

    p1 = tf.add_paragraph()
    p1.text = "Current vs. Proposed Architecture & Scalability Blueprint"
    p1.font.size = Pt(22)
    p1.font.color.rgb = ACCENT_TEAL
    p1.space_before = Pt(12)

    p2 = tf.add_paragraph()
    p2.text = "A technical breakdown of scaling from a single-node prototype to a resilient, distributed, E2EE-ready architecture."
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(18)

    # Footer info
    ftbox = slide1.shapes.add_textbox(Inches(1.2), Inches(6.0), Inches(10), Inches(0.8))
    ft_tf = ftbox.text_frame
    fp = ft_tf.paragraphs[0]
    fp.text = "Prepared for System Architecture Presentation | August 2026"
    fp.font.size = Pt(12)
    fp.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 2: Executive Summary
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, DARK_BG)
    add_header(slide2, "Executive Summary: Evolution of SeChat")

    # Card 1: Problem
    card1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD_BG
    card1.line.color.rgb = BORDER_COLOR
    tf1 = card1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.3)
    tf1.margin_right = Inches(0.3)
    tf1.margin_top = Inches(0.3)

    p = tf1.paragraphs[0]
    p.text = "⚠️ Current Limits"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(248, 113, 113)

    bullets1 = [
        "Monolithic FastAPI process with in-memory connection manager.",
        "SQLite file locks under concurrent write operations.",
        "Cannot scale beyond a single server instance.",
        "Loss of connections during network hiccups."
    ]
    for b in bullets1:
        bp = tf1.add_paragraph()
        bp.text = "• " + b
        bp.font.size = Pt(13)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_before = Pt(10)

    # Card 2: Proposed Solution
    card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.86), Inches(1.8), Inches(3.6), Inches(4.8))
    card2.fill.solid()
    card2.fill.fore_color.rgb = CARD_BG
    card2.line.color.rgb = BORDER_COLOR
    tf2 = card2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.3)
    tf2.margin_right = Inches(0.3)
    tf2.margin_top = Inches(0.3)

    p = tf2.paragraphs[0]
    p.text = "🚀 Proposed System"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL

    bullets2 = [
        "Decoupled Pub/Sub Broker for cross-process broadcasting.",
        "Central Connection Registry/Pool to support multi-worker clusters.",
        "Integrated Crypto Layer (AES-GCM + KDF) for security.",
        "Monitoring & Rate Limiting guardrails for production stability."
    ]
    for b in bullets2:
        bp = tf2.add_paragraph()
        bp.text = "• " + b
        bp.font.size = Pt(13)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_before = Pt(10)

    # Card 3: Business Impact
    card3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.93), Inches(1.8), Inches(3.6), Inches(4.8))
    card3.fill.solid()
    card3.fill.fore_color.rgb = CARD_BG
    card3.line.color.rgb = BORDER_COLOR
    tf3 = card3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = Inches(0.3)
    tf3.margin_right = Inches(0.3)
    tf3.margin_top = Inches(0.3)

    p = tf3.paragraphs[0]
    p.text = "📈 Key Impact"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    bullets3 = [
        "Infinite Horizontal Scalability (Add WS nodes dynamically).",
        "High Availability & Reconnect Resilience.",
        "Enterprise-grade Security (E2EE/MLS roadmap).",
        "Protection against DDoS and message flooding."
    ]
    for b in bullets3:
        bp = tf3.add_paragraph()
        bp.text = "• " + b
        bp.font.size = Pt(13)
        bp.font.color.rgb = TEXT_LIGHT
        bp.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 3: Current Architecture Deep Dive
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, DARK_BG)
    add_header(slide3, "Current System Architecture: Baseline Model")

    # Diagram Representation Box
    diag_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.73), Inches(2.2))
    diag_box.fill.solid()
    diag_box.fill.fore_color.rgb = CARD_BG
    diag_box.line.color.rgb = BORDER_COLOR

    # Draw horizontal sequence boxes inside diagram
    steps = [
        "Browser Client",
        "HTTP REST (FastAPI)",
        "Token Check",
        "WebSocket Server",
        "Connection Manager",
        "SQLite"
    ]
    box_w = Inches(1.6)
    spacing = Inches(0.3)
    start_x = Inches(1.0)
    y_pos = Inches(2.4)

    for idx, step in enumerate(steps):
        s_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, start_x + idx * (box_w + spacing), y_pos, box_w, Inches(1.0))
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = RGBColor(15, 23, 42)
        s_box.line.color.rgb = ACCENT_TEAL
        stf = s_box.text_frame
        stf.word_wrap = True
        sp = stf.paragraphs[0]
        sp.text = step
        sp.font.size = Pt(12)
        sp.font.bold = True
        sp.font.color.rgb = TEXT_LIGHT
        sp.alignment = PP_ALIGN.CENTER

    # Detailed Points below
    info_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.73), Inches(2.4))
    info_box.fill.solid()
    info_box.fill.fore_color.rgb = CARD_BG
    info_box.line.color.rgb = BORDER_COLOR
    itf = info_box.text_frame
    itf.word_wrap = True
    itf.margin_left = Inches(0.4)
    itf.margin_top = Inches(0.3)

    ip = itf.paragraphs[0]
    ip.text = "Key Bottlenecks of Current Design:"
    ip.font.size = Pt(16)
    ip.font.bold = True
    ip.font.color.rgb = RGBColor(248, 113, 113)

    c_bullets = [
        "Tightly Coupled Monolith: HTTP REST, authentication, and WebSocket handling run in a single process.",
        "In-Memory State Lock: Connection Manager maintains active connections strictly in Python memory. Scaling to 2 processes breaks message delivery.",
        "Database Concurrency Bottleneck: Synchronous writes to SQLite cause database lock contention during concurrent chat spikes."
    ]
    for cb in c_bullets:
        p = itf.add_paragraph()
        p.text = "• " + cb
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 4: Proposed Architecture Deep Dive
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, DARK_BG)
    add_header(slide4, "Proposed System Architecture: Distributed & Decoupled")

    # Left Column: Structure Overview
    l_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9))
    l_box.fill.solid()
    l_box.fill.fore_color.rgb = CARD_BG
    l_box.line.color.rgb = BORDER_COLOR
    ltf = l_box.text_frame
    ltf.word_wrap = True
    ltf.margin_left = Inches(0.3)
    ltf.margin_top = Inches(0.3)

    lp = ltf.paragraphs[0]
    lp.text = "Architecture Components"
    lp.font.size = Pt(18)
    lp.font.bold = True
    lp.font.color.rgb = ACCENT_TEAL

    p_nodes = [
        ("Clients", "Web/Mobile connecting via REST & WS Gateways."),
        ("REST API (Auth + Room)", "Dedicated stateless service for user auth & room metadata."),
        ("WebSocket Gateway", "Lightweight edge servers handling full-duplex TCP framing."),
        ("Connection Pool / Registry", "Centralized tracker for active connections & worker routing."),
        ("Pub/Sub Broker", "Redis-like event backbone for cross-worker message fanout."),
        ("Crypto Layer (AES-GCM + KDF)", "E2EE payload encryption & key derivation engine."),
        ("Persistence & Monitoring", "SQL DB persistence paired with real-time rate limiting.")
    ]
    for name, desc in p_nodes:
        p = ltf.add_paragraph()
        p.text = f"• {name}: {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(6)

    # Right Column: Architectural Enhancements
    r_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.9))
    r_box.fill.solid()
    r_box.fill.fore_color.rgb = CARD_BG
    r_box.line.color.rgb = BORDER_COLOR
    rtf = r_box.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.3)
    rtf.margin_top = Inches(0.3)

    rp = rtf.paragraphs[0]
    rp.text = "Core Architecture Enhancements"
    rp.font.size = Pt(18)
    rp.font.bold = True
    rp.font.color.rgb = ACCENT_BLUE

    enhancements = [
        "Process Decoupling: Redis-like Pub/Sub decouples message dissemination from individual WebSocket process memory.",
        "Multi-Worker Connection Registry: Registry supports distributed worker pools and transparent reconnect handling.",
        "Application-Layer Multicast: Maintains logical group delivery without broadcasting overhead across irrelevant nodes.",
        "Advanced Key Management: Cryptographic foundation designed to evolve smoothly toward MLS / Ratcheting protocols."
    ]
    for enh in enhancements:
        p = rtf.add_paragraph()
        p.text = "✔ " + enh
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 5: Head-to-Head Comparison Matrix
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, DARK_BG)
    add_header(slide5, "Head-to-Head Comparison: Baseline vs. Proposed")

    # Table creation
    rows = 6
    cols = 3
    left = Inches(0.8)
    top = Inches(1.8)
    width = Inches(11.73)
    height = Inches(4.8)

    table_shape = slide5.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(2.73)
    table.columns[1].width = Inches(4.5)
    table.columns[2].width = Inches(4.5)

    headers = ["Dimension", "Current System (Baseline)", "Proposed System (Target)"]
    matrix_data = [
        ["Architecture Pattern", "Single-Process Monolith", "Decoupled Micro-services / Gateways"],
        ["Scaling Model", "Vertical only (limited by single process)", "Horizontal (Add WS worker nodes + Pub/Sub)"],
        ["State Management", "In-memory Python ConnectionManager", "Centralized Connection Registry + Pub/Sub"],
        ["Security & Cryptography", "Transport-layer only (HTTPS/WSS)", "Payload Encryption (AES-GCM + KDF / MLS)"],
        ["Fault Tolerance & DDoS", "No rate limits; connection drop risks data loss", "Rate Limiting, Monitoring & Auto-Reconnect"]
    ]

    # Style Header Row
    for i, head in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.text = head
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = ACCENT_TEAL

    # Style Data Rows
    for r_idx, row_data in enumerate(matrix_data):
        for c_idx, text in enumerate(row_data):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            if c_idx == 0:
                cell.fill.fore_color.rgb = RGBColor(30, 41, 59)
            elif c_idx == 1:
                cell.fill.fore_color.rgb = RGBColor(15, 23, 42)
            else:
                cell.fill.fore_color.rgb = RGBColor(20, 35, 55)

            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_LIGHT

    # -------------------------------------------------------------
    # SLIDE 6: Key Technological Advantages
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, DARK_BG)
    add_header(slide6, "Key Technological Advantages & Business Value")

    grid_items = [
        ("1. High Throughput & Low Latency", "Pub/Sub architecture allows instant message routing across sub-millisecond event buses, keeping latency < 20ms even under heavy load.", ACCENT_TEAL),
        ("2. Zero Downtime Horizontal Scaling", "Spin up or tear down WebSocket Gateway containers dynamically based on traffic demand without disrupting active user sessions.", ACCENT_BLUE),
        ("3. End-to-End Cryptographic Trust", "AES-GCM with Key Derivation Functions ensures that data payloads remain encrypted in transit and at rest.", RGBColor(168, 85, 247)),
        ("4. Operational Observability", "Built-in rate limiting prevents malicious spam, while monitoring tools deliver live insights into server health and connection metrics.", RGBColor(245, 158, 11))
    ]

    coords = [
        (Inches(0.8), Inches(1.8)),
        (Inches(6.8), Inches(1.8)),
        (Inches(0.8), Inches(4.3)),
        (Inches(6.8), Inches(4.3))
    ]

    for idx, (title, desc, color) in enumerate(grid_items):
        x, y = coords[idx]
        gbox = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(2.3))
        gbox.fill.solid()
        gbox.fill.fore_color.rgb = CARD_BG
        gbox.line.color.rgb = BORDER_COLOR
        gtf = gbox.text_frame
        gtf.word_wrap = True
        gtf.margin_left = Inches(0.3)
        gtf.margin_top = Inches(0.3)

        p = gtf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = color

        p2 = gtf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 7: Conclusion & Summary
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, DARK_BG)
    add_header(slide7, "Conclusion & Next Steps")

    c_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.73), Inches(4.9))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = CARD_BG
    c_box.line.color.rgb = BORDER_COLOR
    ctf = c_box.text_frame
    ctf.word_wrap = True
    ctf.margin_left = Inches(0.5)
    ctf.margin_top = Inches(0.4)

    cp = ctf.paragraphs[0]
    cp.text = "Summary of Architecture Value Proposition:"
    cp.font.size = Pt(20)
    cp.font.bold = True
    cp.font.color.rgb = ACCENT_TEAL

    c_points = [
        "Eliminates Single Point of Failure: Distributed WebSocket gateways backed by Redis Pub/Sub ensure uninterrupted service.",
        "Future-Proof Cryptography: Integrated AES-GCM + KDF layer provides immediate payload security and sets up Double Ratchet / MLS integration.",
        "Production Readiness: Rate limiting and monitoring safeguards system integrity against spikes and unauthorized traffic.",
        "Seamless User Experience: Connection pool handles silent reconnects, eliminating frustrating user drops."
    ]

    for pt in c_points:
        p = ctf.add_paragraph()
        p.text = "✔ " + pt
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(14)

    # Final Q&A banner
    qa_p = ctf.add_paragraph()
    qa_p.text = "\nThank You! Open for Questions & Technical Discussion."
    qa_p.font.size = Pt(16)
    qa_p.font.bold = True
    qa_p.font.color.rgb = ACCENT_BLUE

    prs.save("c:/Users/vallavha/Desktop/cn_chat/SeChat_Architecture_Presentation.pptx")
    print("Successfully generated SeChat_Architecture_Presentation.pptx")

if __name__ == "__main__":
    create_presentation()
